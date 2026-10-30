#!/usr/bin/env python3
"""Option block trades on J-NET, matched to the day's OI change, with history.

J-NET is where institutions cross option blocks off the book. For each large
contract traded there this records who took part and classifies the trade:

  * 単独クロス — one participant on both sides (its volume is twice the lots):
    a cross inside one broker, between its clients or against its own book;
  * 相対 — the largest participant on one side and the rest on the other,
    when the rest add up to roughly the same lots;
  * 複数 — several participants that do not pair up cleanly.

The lots are then held against the contract's OI change to say whether the
block left new open interest (新規建て), closed it (決済), or only moved it
between holders (付け替え). Participant volume has no buy/sell side, so who
bought and who sold is never inferred.

Output data/opt_blocks.json:
    {"dates": [...], "days": {date: {"blocks": [...], "top_volume": [...]}}}
"""
import argparse
import io
import json
import os
import re
import subprocess

import openpyxl

from extract_jnet import _classify

MIN_LOTS = 50        # smaller J-NET prints are not blocks
TINY_LEG = 0.05      # legs under 5% of the largest, or under 10 lots, are
MIN_LEG = 10         # noise next to a block (retail SBI 2-6 lots, ...)
PAIR_TOL = 0.15      # rest within ±15% of the top leg -> the two sides pair up
KEEP = 30
_CONTRACT = re.compile(r'([PC])(\d{4})-(\d+)$')
JNET_FILES = ('_volume_by_participant_whole_day_J-NET.xlsx', '_volume_by_participant_night_J-NET.xlsx')


def _wb(src):
    return openpyxl.load_workbook(io.BytesIO(src) if isinstance(src, bytes) else src,
                                  read_only=True, data_only=True)


def parse_oi(src):
    """{(side, 'YYMM', strike): (volume, oi, oi_change)} for monthly large options."""
    out = {}
    wb = _wb(src)
    ws = wb['別紙1'] if '別紙1' in wb.sheetnames else wb.worksheets[0]
    for r in ws.iter_rows(values_only=True):
        # Put and call blocks sit side by side, but the call block does not
        # start at a fixed column, so look for contract names anywhere.
        for i in range(len(r)):
            if len(r) < i + 4 or not isinstance(r[i], str):
                continue
            m = _CONTRACT.search(str(r[i]).strip())
            if not m or 'NIKKEI 225' not in str(r[i]):
                continue
            try:
                out[(m.group(1), m.group(2), int(m.group(3)))] = (
                    int(r[i + 1] or 0), int(r[i + 2] or 0), int(r[i + 3] or 0))
            except (TypeError, ValueError):
                pass
    return out


def parse_jnet(src, session):
    """{(side, 'YYMM', strike): {broker: [volume, {session}]}} for option rows."""
    out = {}
    for r in _wb(src).worksheets[0].iter_rows(values_only=True):
        r = [c for c in r if c is not None]
        if len(r) < 7 or str(r[0]).strip() != 'NK225E':
            continue
        m = _CONTRACT.search(str(r[2]).strip())
        if not m:
            continue
        try:
            vol = float(r[-1])
        except (TypeError, ValueError):
            continue
        cell = out.setdefault((m.group(1), m.group(2), int(m.group(3))), {}).setdefault(
            str(r[5]).strip(), [0.0, set()])
        cell[0] += vol
        cell[1].add(session)
    return out


def _merge(a, b):
    for k, brokers in b.items():
        dst = a.setdefault(k, {})
        for name, (v, ss) in brokers.items():
            cell = dst.setdefault(name, [0.0, set()])
            cell[0] += v
            cell[1] |= ss
    return a


def classify(legs):
    """legs sorted by volume desc -> (kind, lots, side_a, side_b)."""
    top = legs[0]['vol']
    sig = [x for x in legs if x['vol'] >= max(MIN_LEG, top * TINY_LEG)]
    if len(sig) == 1:
        return 'single', top / 2.0, sig[:1], []
    rest = sum(x['vol'] for x in sig[1:])
    if abs(rest - top) <= top * PAIR_TOL:
        return 'pair', (top + rest) / 2.0, sig[:1], sig[1:]
    return 'multi', sum(x['vol'] for x in sig) / 2.0, sig, []


def verdict(lots, chg):
    if chg is None or not lots:
        return None
    r = chg / lots
    if r >= 0.7:
        return '新規建て'
    if r <= -0.7:
        return '決済'
    if abs(r) <= 0.3:
        return '付け替え'
    return '新規と決済が混在'


def build_day(oi, jnet):
    blocks = []
    for (side, exp, k), brokers in jnet.items():
        legs = sorted(({'broker': b, 'vol': int(v), 'cat': _classify(b), 'sessions': sorted(ss)}
                       for b, (v, ss) in brokers.items()), key=lambda x: -x['vol'])
        kind, lots, a, b = classify(legs)
        if lots < MIN_LOTS:
            continue
        vol, cur, chg = oi.get((side, exp, k), (None, None, None))
        sessions = sorted({s for x in legs for s in x['sessions']})
        blocks.append({'side': side, 'exp': exp, 'k': k, 'lots': int(round(lots)), 'kind': kind,
                       'a': a, 'b': b, 'sessions': sessions,
                       'vol': vol, 'oi': cur, 'chg': chg, 'verdict': verdict(lots, chg)})
    blocks.sort(key=lambda x: -x['lots'])
    top_volume = sorted(({'side': s, 'exp': e, 'k': k, 'vol': v, 'oi': o, 'chg': c}
                         for (s, e, k), (v, o, c) in oi.items() if v), key=lambda x: -x['vol'])[:6]
    return {'blocks': blocks, 'top_volume': top_volume}


def _load(path):
    if os.path.exists(path):
        try:
            with open(path, encoding='utf-8') as f:
                d = json.load(f)
            if isinstance(d.get('days'), dict):
                return d
        except Exception:
            pass
    return {'dates': [], 'days': {}}


def _git_files(repo_root, suffix):
    """{date: (sha, path)} for files with this suffix ever added to data/."""
    out = {}
    try:
        log = subprocess.check_output(['git', '-C', repo_root, 'log', '--diff-filter=A', '--format=%H',
                                       '--name-only', '--', 'data/*' + suffix], text=True)
    except (OSError, subprocess.CalledProcessError):
        return out
    sha = None
    for ln in log.splitlines():
        ln = ln.strip()
        if re.fullmatch(r'[0-9a-f]{40}', ln):
            sha = ln
        elif sha:
            m = re.search(r'(\d{8})' + re.escape(suffix) + '$', ln)
            if m:
                out.setdefault(m.group(1), (sha, ln))
    return out


def _backfill_git(store, repo_root):
    show = lambda sp: subprocess.check_output(['git', '-C', repo_root, 'show', '%s:%s' % sp])
    ois = _git_files(repo_root, 'open_interest.xlsx')
    day = _git_files(repo_root, JNET_FILES[0])
    night = _git_files(repo_root, JNET_FILES[1])
    for dt in sorted(set(ois) & set(day)):
        if dt in store['days']:
            continue
        try:
            jn = parse_jnet(show(day[dt]), 'day')
            if dt in night:
                _merge(jn, parse_jnet(show(night[dt]), 'night'))
            store['days'][dt] = build_day(parse_oi(show(ois[dt])), jn)
        except Exception as e:
            print('[opt_blocks] backfill %s failed: %s' % (dt, e))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--data-dir', default='data')
    ap.add_argument('--out', default=None, help='default: <data-dir>/opt_blocks.json')
    ap.add_argument('--backfill-git', action='store_true')
    a = ap.parse_args()
    out = a.out or os.path.join(a.data_dir, 'opt_blocks.json')
    store = _load(out)

    oi_dates = [m.group(1) for m in (re.fullmatch(r'(\d{8})open_interest\.xlsx', f) for f in os.listdir(a.data_dir)) if m]
    for dt in sorted(oi_dates)[-1:]:
        day = os.path.join(a.data_dir, dt + JNET_FILES[0])
        if not os.path.exists(day):
            print('[opt_blocks] %s: day J-NET file missing; history unchanged' % dt)
            continue
        night = os.path.join(a.data_dir, dt + JNET_FILES[1])
        if not os.path.exists(night) and dt in store['days']:
            # Re-runs after the workflow's cleanup see only the day file; keep
            # the entry that was built with both sessions.
            print('[opt_blocks] %s: night J-NET file missing; keeping stored entry' % dt)
            continue
        jn = parse_jnet(day, 'day')
        if os.path.exists(night):
            _merge(jn, parse_jnet(night, 'night'))
        store['days'][dt] = build_day(parse_oi(os.path.join(a.data_dir, dt + 'open_interest.xlsx')), jn)

    if a.backfill_git:
        _backfill_git(store, os.path.dirname(os.path.abspath(a.data_dir)))

    keep = sorted(store['days'])[-KEEP:]
    store = {'dates': keep, 'days': {d: store['days'][d] for d in keep}}
    with open(out, 'w', encoding='utf-8') as f:
        json.dump(store, f, ensure_ascii=False, separators=(',', ':'))
    last = store['days'].get(keep[-1], {}) if keep else {}
    print('[opt_blocks] wrote %s | %d days | latest %s: %d blocks'
          % (out, len(keep), keep[-1] if keep else '-', len(last.get('blocks', []))))


if __name__ == '__main__':
    main()
