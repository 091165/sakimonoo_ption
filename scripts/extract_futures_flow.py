#!/usr/bin/env python3
"""Daily futures flow history for the 先物 card: price x open interest x volume.

A raw OI change ("-2,346") says little on its own. Read together with the
day's price move it tells who acted:

    price up   + OI up   -> 新規買い優勢   (new longs)
    price up   + OI down -> 売り方の買い戻し (short covering)
    price down + OI up   -> 新規売り優勢   (new shorts)
    price down + OI down -> 買い方の手仕舞い (long liquidation)

Price is the Nikkei 225 cash close taken from ose<date>tp.csv (the option
underlying), which is the same number every day regardless of how the
pipeline was run. OI comes from data.json (s02) for today and from
oi_timeseries.json for past days; volume (day + night, all sessions) from
<date>_derivatives_market_data_whole_day.xlsx.

The output keeps its own history (upsert by date), because the daily market
data file is deleted from data/ after each run. --backfill-git rebuilds past
volumes once from the files still reachable in git history.

Output data/futures_flow.json:
    {"dates": [...], "rows": {date: {"close": float|None,
        "nk225_large"|"nk225_mini"|"topix": {"oi", "chg", "vol"}}}}
"""
import argparse
import glob
import io
import json
import os
import re
import subprocess

import openpyxl

PRODUCTS = ('nk225_large', 'nk225_mini', 'topix')
# market data row label prefix -> product key
_VOL_LABELS = (('日経225先物', 'nk225_large'), ('日経225mini', 'nk225_mini'), ('TOPIX', 'topix'))
KEEP = 60


def _close_from_tp(path):
    """Nikkei 225 cash close = underlying column of the front NK225E rows."""
    try:
        with open(path, encoding='utf-8', errors='replace') as f:
            for ln in f:
                p = ln.split(',')
                if len(p) > 15 and p[0].strip() == 'NK225E':
                    v = float(p[15])
                    if v > 1000:
                        return round(v, 2)
    except (OSError, ValueError):
        pass
    return None


def volumes_from_market_data(src):
    """{product: total volume} from the futures market data sheet (path or bytes)."""
    wb = openpyxl.load_workbook(io.BytesIO(src) if isinstance(src, bytes) else src,
                                read_only=True, data_only=True)
    name = next((n for n in wb.sheetnames if n in ('market_data_Futures', '指数先物')), None)
    if not name:
        return {}
    out, cur = {}, None
    for r in wb[name].iter_rows(values_only=True):
        r = list(r) + [None] * 6
        label = str(r[2] or '').replace('\n', ' ').strip()
        if label:
            cur = next((k for pre, k in _VOL_LABELS if label.startswith(pre)), None)
        session = str(r[3] or '')
        if cur and cur not in out and ('合計' in session or 'Total' in session):
            try:
                out[cur] = int(float(r[4]))
            except (TypeError, ValueError):
                pass
    return out


def _load(path):
    if os.path.exists(path):
        try:
            with open(path, encoding='utf-8') as f:
                d = json.load(f)
            if isinstance(d.get('rows'), dict):
                return d
        except Exception:
            pass
    return {'dates': [], 'rows': {}}


def _seed_from_history(flow, data_dir):
    """Fill OI (and closes) for past days from oi_timeseries.json + tp.csv files."""
    path = os.path.join(data_dir, 'oi_timeseries.json')
    if os.path.exists(path):
        with open(path, encoding='utf-8') as f:
            ts = json.load(f)
        dates = ts.get('dates') or []
        fut = ts.get('futures') or {}
        for key in PRODUCTS:
            tot = (fut.get(key) or {}).get('total') or []
            for i, dt in enumerate(dates):
                if i >= len(tot) or not tot[i]:
                    continue
                row = flow['rows'].setdefault(dt, {})
                cell = row.setdefault(key, {})
                cell.setdefault('oi', int(tot[i]))
                if i > 0 and tot[i - 1]:
                    cell.setdefault('chg', int(tot[i] - tot[i - 1]))
    for tp in glob.glob(os.path.join(data_dir, 'ose*tp.csv')):
        m = re.search(r'ose(\d{8})tp\.csv$', tp)
        if m and m.group(1) in flow['rows'] and not flow['rows'][m.group(1)].get('close'):
            flow['rows'][m.group(1)]['close'] = _close_from_tp(tp)
    # Days without a tp.csv (e.g. 10/08): fall back to the fetched OHLC history,
    # whose closes match the tp.csv underlying on the days both exist.
    ml = os.path.join(data_dir, 'market_latest.json')
    if os.path.exists(ml):
        with open(ml, encoding='utf-8') as f:
            for bar in (json.load(f).get('ohlc_history') or []):
                dt = str(bar.get('date', '')).replace('-', '')
                if dt in flow['rows'] and not flow['rows'][dt].get('close') and bar.get('close'):
                    flow['rows'][dt]['close'] = round(float(bar['close']), 2)


def _backfill_git(flow, repo_root):
    """One-off: past daily volumes from market data files still in git history."""
    try:
        log = subprocess.check_output(
            ['git', '-C', repo_root, 'log', '--diff-filter=A', '--format=%H', '--name-only',
             '--', 'data/*_derivatives_market_data_whole_day.xlsx'], text=True)
    except (OSError, subprocess.CalledProcessError) as e:
        print('[futures_flow] git backfill skipped: %s' % e)
        return
    sha = None
    for ln in log.splitlines():
        ln = ln.strip()
        if re.fullmatch(r'[0-9a-f]{40}', ln):
            sha = ln
            continue
        m = re.search(r'(\d{8})_derivatives_market_data_whole_day\.xlsx$', ln)
        if not (sha and m and m.group(1) in flow['rows']):
            continue
        try:
            blob = subprocess.check_output(['git', '-C', repo_root, 'show', '%s:%s' % (sha, ln)])
            for key, v in volumes_from_market_data(blob).items():
                flow['rows'][m.group(1)].setdefault(key, {}).setdefault('vol', v)
        except Exception as e:
            print('[futures_flow] backfill %s failed: %s' % (ln, e))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--data-dir', default='data')
    ap.add_argument('--data', default=None, help='data.json (default: <data-dir>/data.json)')
    ap.add_argument('--out', default=None, help='default: <data-dir>/futures_flow.json')
    ap.add_argument('--backfill-git', action='store_true')
    a = ap.parse_args()
    data_path = a.data or os.path.join(a.data_dir, 'data.json')
    out = a.out or os.path.join(a.data_dir, 'futures_flow.json')

    flow = _load(out)
    _seed_from_history(flow, a.data_dir)

    with open(data_path, encoding='utf-8') as f:
        d = json.load(f)
    today = (d.get('metadata') or {}).get('date')
    s02 = d.get('s02') or {}
    if today and 'error' not in s02:
        row = flow['rows'].setdefault(today, {})
        for key in PRODUCTS:
            sec = s02.get(key) or {}
            if sec.get('total_oi'):
                cell = row.setdefault(key, {})
                cell['oi'] = int(sec['total_oi'])
                cell['chg'] = int(sec.get('total_change') or 0)
        row['close'] = _close_from_tp(os.path.join(a.data_dir, 'ose%stp.csv' % today)) or row.get('close')
        md = os.path.join(a.data_dir, '%s_derivatives_market_data_whole_day.xlsx' % today)
        if os.path.exists(md):
            for key, v in volumes_from_market_data(md).items():
                row.setdefault(key, {})['vol'] = v

    if a.backfill_git:
        _backfill_git(flow, os.path.dirname(os.path.abspath(a.data_dir)))

    dates = sorted(flow['rows'])[-KEEP:]
    flow = {'dates': dates, 'rows': {dt: flow['rows'][dt] for dt in dates}}
    with open(out, 'w', encoding='utf-8') as f:
        json.dump(flow, f, ensure_ascii=False, indent=1)
    last = flow['rows'].get(dates[-1], {}) if dates else {}
    print('[futures_flow] wrote %s | %d days | latest %s close=%s large=%s'
          % (out, len(dates), dates[-1] if dates else '-', last.get('close'), last.get('nk225_large')))


if __name__ == '__main__':
    main()
